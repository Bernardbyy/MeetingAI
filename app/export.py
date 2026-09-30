"""minutes.md -> a Word document people can edit and send on.

The minutes are the four-section Markdown summarize.py asks for (headings,
bullets, paragraphs, **bold**), read line by line the same way the viewer
renders them. Word's own styles are used (Heading, List Bullet), so the
document restyles and navigates like one typed in Word.
"""

import io
import re

from docx import Document


def _add_text(par, text: str):
    """Append text to a paragraph, turning **bold** into bold runs."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            par.add_run(part).bold = i % 2 == 1


def minutes_docx(minutes: str, title: str, subtitle: str = "") -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    if subtitle:
        doc.add_paragraph(subtitle).runs[0].italic = True

    for raw in minutes.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("#"):
            doc.add_heading(line.lstrip("#").strip(), level=2)
        elif re.match(r"^[-*]\s+", line):
            _add_text(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*]\s+", "", line))
        elif line.startswith(">"):
            par = doc.add_paragraph()
            _add_text(par, line.lstrip("> "))
            for run in par.runs:
                run.italic = True
        else:
            _add_text(doc.add_paragraph(), line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
